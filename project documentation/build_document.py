from pathlib import Path
import re
import textwrap
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.shared import Mm
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image as PdfImage
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from xml.sax.saxutils import escape


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "Project Documentation.md"
WORD_OUTPUT = ROOT / "Emergency Unit Requisition Platform - Project Documentation.docx"
PDF_OUTPUT = ROOT / "Emergency Unit Requisition Platform - Project Documentation.pdf"
REGULAR_FONT = Path(r"C:\Windows\Fonts\arial.ttf")
BOLD_FONT = Path(r"C:\Windows\Fonts\arialbd.ttf")


def inline_markup(text):
    text = escape(text)
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"\*(.+?)\*", r"<i>\1</i>", text)
    text = re.sub(r"`(.+?)`", r"<font name='Courier'>\1</font>", text)
    return text


def parse_table(lines):
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        rows.append(cells)
    return rows


def add_inline_runs(paragraph, text):
    tokens = re.split(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)", text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("*") and token.endswith("*"):
            paragraph.add_run(token[1:-1]).italic = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(token)


def parse_blocks(section):
    lines = section.strip().splitlines()
    blocks = []
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            blocks.append(("table", parse_table(table_lines)))
            continue
        image_match = re.fullmatch(r"!\[(.+?)\]\((.+?)\)", line)
        if image_match:
            blocks.append(("image", image_match.group(1), image_match.group(2)))
            index += 1
            continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            blocks.append(("heading", level, line[level:].strip()))
            index += 1
            continue
        if line.startswith("- "):
            blocks.append(("bullet", line[2:].strip()))
            index += 1
            continue
        if re.match(r"^\d+\.\s", line):
            blocks.append(("paragraph", line))
            index += 1
            continue
        paragraph = [line]
        index += 1
        while index < len(lines):
            next_line = lines[index].strip()
            if not next_line or next_line.startswith(("#", "|", "- ")) or re.match(r"^\d+\.\s", next_line):
                break
            paragraph.append(next_line)
            index += 1
        blocks.append(("paragraph", " ".join(paragraph)))
    return blocks


def add_word_table(document, rows):
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=1, cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, row in enumerate(rows):
        cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
        for column in range(columns):
            cell = cells[column]
            cell.text = ""
            value = row[column] if column < len(row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(1)
            run = paragraph.add_run(value)
            run.font.name = "Aptos"
            run.font.size = Pt(8)
            if row_index == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "146D68")
                cell._tc.get_or_add_tcPr().append(shading)
            elif row_index % 2 == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "F0F6F4")
                cell._tc.get_or_add_tcPr().append(shading)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    document.add_paragraph()


def create_diagrams():
    regular = ImageFont.truetype(str(REGULAR_FONT), 23) if REGULAR_FONT.exists() else ImageFont.load_default()
    bold = ImageFont.truetype(str(BOLD_FONT), 28) if BOLD_FONT.exists() else ImageFont.load_default()
    title_font = ImageFont.truetype(str(BOLD_FONT), 36) if BOLD_FONT.exists() else ImageFont.load_default()
    ink = "#183C42"
    green = "#146D68"
    pale = "#EAF4F1"
    white = "#FFFFFF"
    muted = "#527275"

    def centered_lines(draw, center, text, font, width, fill):
        words = text.split()
        lines = []
        current = ""
        for word in words:
            candidate = f"{current} {word}".strip()
            if current and draw.textlength(candidate, font=font) > width:
                lines.append(current)
                current = word
            else:
                current = candidate
        if current:
            lines.append(current)
        heights = [draw.textbbox((0, 0), line, font=font)[3] for line in lines]
        y = center[1] - sum(heights) / 2 - 4 * (len(lines) - 1)
        for line, height in zip(lines, heights):
            draw.text((center[0] - draw.textlength(line, font=font) / 2, y), line, font=font, fill=fill)
            y += height + 8

    def arrow(draw, start, end, color=green, width=7):
        draw.line([start, end], fill=color, width=width)
        angle = __import__("math").atan2(end[1] - start[1], end[0] - start[0])
        size = 22
        left = (end[0] - size * __import__("math").cos(angle - 0.5), end[1] - size * __import__("math").sin(angle - 0.5))
        right = (end[0] - size * __import__("math").cos(angle + 0.5), end[1] - size * __import__("math").sin(angle + 0.5))
        draw.polygon([end, left, right], fill=color)

    image = Image.new("RGB", (1700, 560), "#F7FAF9")
    draw = ImageDraw.Draw(image)
    draw.text((55, 32), "SYSTEM ARCHITECTURE", fill=ink, font=title_font)
    nodes = [
        ("Browser", "React pages and role-specific workspaces"),
        ("Vite proxy", "Forwards /api requests during local development"),
        ("PHP API", "Single entry point, routing, authentication and permissions"),
        ("Application", "Controllers and models validate and process requests"),
        ("Database", "PDO connects to SQLite or configured MySQL"),
    ]
    x_positions = [55, 390, 725, 1060, 1395]
    for index in range(len(nodes) - 1):
        arrow(draw, (x_positions[index] + 245, 290), (x_positions[index + 1] - 18, 290))
    for x, (heading, detail) in zip(x_positions, nodes):
        draw.rounded_rectangle((x, 165, x + 245, 410), radius=25, fill=white, outline=green, width=5)
        centered_lines(draw, (x + 122, 220), heading, bold, 205, green)
        centered_lines(draw, (x + 122, 315), detail, regular, 205, ink)
    centered_lines(draw, (850, 485), "Protected requests carry a signed token; the API returns structured JSON.", regular, 1500, muted)
    image.save(ROOT / "system-architecture.png")

    image = Image.new("RGB", (1500, 970), "#F7FAF9")
    draw = ImageDraw.Draw(image)
    draw.text((55, 28), "DATABASE RELATIONSHIPS", fill=ink, font=title_font)
    boxes = {
        "users": (55, 155, 405, 480, "Users", ["PK  id", "username", "email", "password_hash", "role", "approval_status"]),
        "req": (575, 155, 965, 530, "Requisitions", ["PK  id", "FK  item_id", "FK  user_id", "quantity", "status", "admin_notes", "created_at / reviewed_at"]),
        "items": (1120, 155, 1460, 480, "Inventory items", ["PK  id", "name", "category", "stock", "status", "is_active"]),
        "activity": (55, 640, 460, 920, "Activity logs", ["PK  id", "FK  actor_user_id", "actor_role", "action", "page", "created_at"]),
    }
    arrow(draw, (405, 320), (575, 320))
    arrow(draw, (965, 320), (1120, 320))
    draw.line([(230, 480), (230, 780), (575, 780)], fill=green, width=7)
    arrow(draw, (460, 780), (555, 780))
    draw.text((415, 280), "1 : many", font=regular, fill=muted)
    draw.text((990, 280), "1 : many", font=regular, fill=muted)
    draw.text((260, 690), "1 : many", font=regular, fill=muted)
    for _, (x1, y1, x2, y2, heading, fields) in boxes.items():
        draw.rounded_rectangle((x1, y1, x2, y2), radius=22, fill=white, outline=green, width=5)
        draw.rounded_rectangle((x1, y1, x2, y1 + 68), radius=20, fill=pale, outline=green, width=3)
        draw.text((x1 + 18, y1 + 17), heading, font=bold, fill=green)
        y = y1 + 92
        for field in fields:
            draw.text((x1 + 22, y), field, font=regular, fill=ink)
            y += 38
    centered_lines(draw, (1010, 785), "An activity record belongs to the account that performed the action.", regular, 460, muted)
    image.save(ROOT / "database-relationships.png")

    image = Image.new("RGB", (1600, 670), "#F7FAF9")
    draw = ImageDraw.Draw(image)
    draw.text((55, 28), "REQUISITION REVIEW FLOW", fill=ink, font=title_font)
    main_nodes = [
        (55, "Unit User", "Chooses an active item and submits a quantity"),
        (425, "API checks", "Confirms role, item and available stock"),
        (795, "Pending", "Saves the request for administrator review"),
        (1165, "Administrator", "Records approval or decline"),
    ]
    for index in range(len(main_nodes) - 1):
        arrow(draw, (main_nodes[index][0] + 275, 240), (main_nodes[index + 1][0] - 20, 240))
    for x, heading, detail in main_nodes:
        draw.rounded_rectangle((x, 140, x + 275, 340), radius=24, fill=white, outline=green, width=5)
        centered_lines(draw, (x + 137, 188), heading, bold, 240, green)
        centered_lines(draw, (x + 137, 270), detail, regular, 235, ink)
    draw.line([(1300, 340), (1300, 405), (610, 405), (610, 445)], fill=green, width=6)
    arrow(draw, (610, 405), (610, 445))
    draw.line([(1300, 405), (1330, 405), (1330, 445)], fill=green, width=6)
    arrow(draw, (1330, 405), (1330, 445))
    draw.rounded_rectangle((420, 445, 800, 615), radius=24, fill="#EAF4F1", outline=green, width=5)
    centered_lines(draw, (610, 490), "Approved", bold, 340, green)
    centered_lines(draw, (610, 560), "Stock is deducted in the review transaction.", regular, 335, ink)
    draw.rounded_rectangle((1115, 445, 1510, 615), radius=24, fill="#FFF3F0", outline="#9C5142", width=5)
    centered_lines(draw, (1312, 490), "Declined", bold, 355, "#8D4235")
    centered_lines(draw, (1312, 560), "Stock is unchanged.", regular, 345, ink)
    image.save(ROOT / "requisition-flow.png")


def build_word(sections):
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(42, 55, 58)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color in (("Title", 28, "124B52"), ("Heading 1", 19, "124B52"), ("Heading 2", 13, "146D68")):
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
        style.paragraph_format.space_before = Pt(7)
        style.paragraph_format.space_after = Pt(5)
    header = section.header.paragraphs[0]
    header.text = "HOSPITAL EMERGENCY REQUISITION PLATFORM  |  PROJECT DOCUMENTATION"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(88, 116, 117)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Project documentation  •  ")
    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    footer._p.append(page_field)

    for section_index, blocks in enumerate(sections):
        if section_index:
            document.add_page_break()
        if section_index == 0:
            title = next((block[2] for block in blocks if block[0] == "heading" and block[1] == 1), "Hospital Emergency Requisition Platform")
            subtitle = next((block[2] for block in blocks if block[0] == "heading" and block[1] == 2), "Project Documentation")
            paragraphs = [block[1] for block in blocks if block[0] == "paragraph"]
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(85)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run("EMERGENCY UNIT  •  DIGITAL WORKFLOW")
            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(20, 109, 104)
            paragraph = document.add_paragraph(style="Title")
            paragraph.paragraph_format.space_before = Pt(30)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run(title)
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(subtitle.upper())
            run.bold = True
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(40, 90, 94)
            for text in paragraphs:
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(18)
                add_inline_runs(paragraph, text)
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(50)
            paragraph.add_run("Prepared as a record of the project design, implementation, and verification")
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(20)
            paragraph.add_run("9 October 2026").bold = True
            continue
        for block in blocks:
            kind = block[0]
            if kind == "heading":
                level, text = block[1], block[2]
                if level == 1:
                    document.add_paragraph(text, style="Heading 1")
                elif level == 2:
                    document.add_paragraph(text, style="Heading 2")
                else:
                    document.add_paragraph(text, style="Heading 3")
            elif kind == "paragraph":
                paragraph = document.add_paragraph()
                add_inline_runs(paragraph, block[1])
            elif kind == "bullet":
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.paragraph_format.space_after = Pt(3)
                add_inline_runs(paragraph, block[1])
            elif kind == "table":
                add_word_table(document, block[1])
            elif kind == "image":
                image_path = ROOT / block[2]
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(str(image_path), width=Inches(6.7))
                caption = document.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption.paragraph_format.space_after = Pt(6)
                run = caption.add_run(block[1])
                run.italic = True
                run.font.size = Pt(8)
    document.core_properties.title = "Hospital Emergency Requisition Platform - Project Documentation"
    document.core_properties.subject = "Project design, implementation, testing, and operating guide"
    document.core_properties.author = "Project Developer"
    document.save(WORD_OUTPUT)


def build_pdf(sections):
    regular_path = REGULAR_FONT
    bold_path = BOLD_FONT
    if regular_path.exists() and bold_path.exists():
        pdfmetrics.registerFont(TTFont("ProjectArial", str(regular_path)))
        pdfmetrics.registerFont(TTFont("ProjectArialBold", str(bold_path)))
        regular, bold = "ProjectArial", "ProjectArialBold"
    else:
        regular, bold = "Helvetica", "Helvetica-Bold"
    sheet = getSampleStyleSheet()
    body = ParagraphStyle("ProjectBody", parent=sheet["BodyText"], fontName=regular, fontSize=9.2, leading=12.1, textColor=colors.HexColor("#2A373A"), spaceAfter=5)
    heading = ParagraphStyle("ProjectHeading", parent=body, fontName=bold, fontSize=17, leading=20, textColor=colors.HexColor("#124B52"), spaceAfter=9, keepWithNext=True)
    subheading = ParagraphStyle("ProjectSubheading", parent=body, fontName=bold, fontSize=11.3, leading=14, textColor=colors.HexColor("#146D68"), spaceBefore=5, spaceAfter=4, keepWithNext=True)
    small = ParagraphStyle("ProjectTable", parent=body, fontSize=7.5, leading=9.1, spaceAfter=1)
    cover = ParagraphStyle("ProjectCover", parent=body, fontName=bold, fontSize=25, leading=31, textColor=colors.HexColor("#124B52"), alignment=TA_CENTER, spaceAfter=10)
    subtitle = ParagraphStyle("ProjectSubtitle", parent=body, fontName=bold, fontSize=15, leading=19, textColor=colors.HexColor("#28605E"), alignment=TA_CENTER)
    story = []
    for section_index, blocks in enumerate(sections):
        if section_index:
            story.append(PageBreak())
        if section_index == 0:
            title = next((block[2] for block in blocks if block[0] == "heading" and block[1] == 1), "Hospital Emergency Requisition Platform")
            sub = next((block[2] for block in blocks if block[0] == "heading" and block[1] == 2), "Project Documentation")
            paragraphs = [block[1] for block in blocks if block[0] == "paragraph"]
            story.extend([
                Spacer(1, 55 * mm),
                Paragraph("EMERGENCY UNIT  •  DIGITAL WORKFLOW", ParagraphStyle("CoverTop", parent=body, fontName=bold, fontSize=9, textColor=colors.HexColor("#146D68"), alignment=TA_CENTER, spaceAfter=14)),
                Paragraph(escape(title), cover),
                Paragraph(escape(sub.upper()), subtitle),
                Spacer(1, 7 * mm),
            ])
            for text in paragraphs:
                story.append(Paragraph(inline_markup(text), ParagraphStyle("CoverBody", parent=body, fontSize=11.2, leading=16, alignment=TA_CENTER)))
            story.extend([Spacer(1, 16 * mm), Paragraph("Prepared as a record of the project design, implementation, and verification", ParagraphStyle("Prepared", parent=body, alignment=TA_CENTER)), Spacer(1, 8 * mm), Paragraph("9 October 2026", ParagraphStyle("Date", parent=body, fontName=bold, alignment=TA_CENTER))])
            continue
        for block in blocks:
            kind = block[0]
            if kind == "heading":
                style = heading if block[1] == 1 else subheading
                story.append(Paragraph(escape(block[2]), style))
            elif kind == "paragraph":
                story.append(Paragraph(inline_markup(block[1]), body))
            elif kind == "bullet":
                story.append(Paragraph("&bull;&nbsp;&nbsp;" + inline_markup(block[1]), body))
            elif kind == "table":
                rows = block[1]
                if not rows:
                    continue
                data = [[Paragraph(inline_markup(cell), small) for cell in row] for row in rows]
                usable_width = A4[0] - 36 * mm
                count = len(rows[0])
                if count == 2:
                    widths = [usable_width * 0.34, usable_width * 0.66]
                elif count == 3:
                    widths = [usable_width * 0.24, usable_width * 0.34, usable_width * 0.42]
                else:
                    widths = [usable_width / count] * count
                table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
                table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#146D68")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), bold),
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#C7D7D5")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0F6F4")]),
                ]))
                story.extend([table, Spacer(1, 4)])
            elif kind == "image":
                image_path = ROOT / block[2]
                with Image.open(image_path) as source_image:
                    aspect = source_image.height / source_image.width
                image = PdfImage(str(image_path), width=174 * mm, height=174 * mm * aspect)
                image.hAlign = "CENTER"
                story.extend([image, Paragraph(escape(block[1]), ParagraphStyle("FigureCaption", parent=body, fontSize=8, leading=10, alignment=TA_CENTER, spaceAfter=7))])

    def decorate(canvas, document):
        canvas.saveState()
        width, height = A4
        canvas.setStrokeColor(colors.HexColor("#C6DEDA"))
        canvas.setLineWidth(0.6)
        canvas.line(18 * mm, height - 15 * mm, width - 18 * mm, height - 15 * mm)
        canvas.setFont(regular, 7.5)
        canvas.setFillColor(colors.HexColor("#587475"))
        canvas.drawRightString(width - 18 * mm, height - 12 * mm, "HOSPITAL EMERGENCY REQUISITION PLATFORM  |  PROJECT DOCUMENTATION")
        canvas.line(18 * mm, 14 * mm, width - 18 * mm, 14 * mm)
        canvas.drawCentredString(width / 2, 9 * mm, f"Project documentation  •  {document.page}")
        canvas.restoreState()

    pdf = SimpleDocTemplate(str(PDF_OUTPUT), pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm, topMargin=22 * mm, bottomMargin=19 * mm, title="Hospital Emergency Requisition Platform - Project Documentation", author="Project Developer")
    pdf.build(story, onFirstPage=decorate, onLaterPages=decorate)


def main():
    if not SOURCE.is_file():
        raise FileNotFoundError(f"Project report source not found: {SOURCE}")
    sections = [parse_blocks(section) for section in SOURCE.read_text(encoding="utf-8").split("\n---\n")]
    create_diagrams()
    build_word(sections)
    build_pdf(sections)
    print(f"Source sections: {len(sections)}")
    print(f"Word report: {WORD_OUTPUT}")
    print(f"PDF report: {PDF_OUTPUT}")


if __name__ == "__main__":
    main()
